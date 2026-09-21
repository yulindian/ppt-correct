"""Create a tiny PPTX with three ordered entrance effects for the audit test."""

import sys

from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches


def main(path: str) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape_ids = []
    for index in range(3):
        shape = slide.shapes.add_textbox(Inches(1), Inches(index + 1), Inches(3), Inches(0.5))
        shape.text = f"Step {index + 1}"
        shape_ids.append(shape.shape_id)

    effects = []
    for index, shape_id in enumerate(shape_ids):
        effects.append(
            f'<p:par><p:cTn id="{index + 3}" presetID="10" presetClass="entr" '
            f'nodeType="clickEffect"><p:childTnLst><p:animEffect transition="in" '
            f'filter="fade"><p:cBhvr><p:cTn id="{index + 6}" dur="350"/>'
            f'<p:tgtEl><p:spTgt spid="{shape_id}"/></p:tgtEl></p:cBhvr>'
            f'</p:animEffect></p:childTnLst></p:cTn></p:par>'
        )
    timing = (
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:tnLst><p:par><p:cTn id="1" dur="indefinite" nodeType="tmRoot">'
        '<p:childTnLst><p:seq><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        f'<p:childTnLst>{"".join(effects)}</p:childTnLst>'
        '</p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )
    slide._element.append(parse_xml(timing))
    presentation.save(path)


if __name__ == "__main__":
    main(sys.argv[1])
